"""
Trial Balance - Group Subtotal Cleaner
======================================
Removes the group-subtotal amounts from a trial balance export so that only
ledger balances remain and Debit total = Credit total.

How it works
------------
1. Finds the header row, the account-name column and the Debit/Credit columns
   (all overridable in the UI).
2. Works out which rows are groups (subtotals) from the hierarchy:
      - leading-space indentation (Tally style), or
      - Excel cell indent, or
      - Excel row-outline level, or
      - amounts only (heuristic fallback when the file carries no hierarchy)
3. VERIFIES every candidate: a group's amount (net of Dr/Cr) must equal the net
   of the ledgers beneath it. Only verified subtotals are pre-selected for
   removal; anything that does not tie is flagged for review.
4. Checks that Debit = Credit after removal, then writes a cleaned workbook
   (original formatting kept) plus a "TB Clean Report" sheet logging every
   amount removed.

Integration notes (for statutory_compliance.py)
-----------------------------------------------
- Everything the tool needs is in `render_tb_subtotal_cleaner()`.
- All session-state / widget keys are prefixed with NS = "tbc_".
- The only standalone boilerplate is the block under `if __name__ == "__main__":`.
- Requires: streamlit>=1.37, openpyxl.
"""
from __future__ import annotations

import csv
import hashlib
import io
import re
from dataclasses import dataclass, field

import streamlit as st
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

NS = "tbc_"
ACCENT = "D85A30"
REPORT_SHEET = "TB Clean Report"

NAME_KEYS = ("account", "particulars", "ledger", "name", "description", "head")
SINGLE_AMT_KEYS = ("closing", "opening", "balance", "amount", "net")
TOTAL_RE = re.compile(r"^\s*(grand\s+)?total\b|^\s*diff(erence)?\b", re.I)
DRCR_SUFFIX_RE = re.compile(r"^(.*?)\s*(dr|cr)\.?$", re.I)
PLAIN_NUM_RE = re.compile(r"[-+]?\d[\d,]*(\.\d+)?")


# ───────────────────────────── data classes ──────────────────────────────
@dataclass
class Layout:
    header_row: int
    name_col: int
    pairs: list = field(default_factory=list)      # [(dr_col, cr_col|None)]
    labels: list = field(default_factory=list)     # label per pair


@dataclass
class Row:
    xl_row: int
    name: str
    lvl_space: int
    lvl_indent: int
    lvl_outline: int
    dr: list
    cr: list
    net: list
    has_amt: bool


# ───────────────────────────── parsing helpers ───────────────────────────
def parse_amount(v):
    """Number from a cell: handles floats, '1,234.50', '(1,234.50)', '1,234.50 Cr'."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if not s or s in {"-", "--", "\u2014"}:
        return None
    sign = 1.0
    m = DRCR_SUFFIX_RE.match(s)
    if m:
        s = m.group(1)
        sign = -1.0 if m.group(2).lower() == "cr" else 1.0
    s = re.sub(r"(?i)rs\.?|inr|\u20b9|[,\s]", "", s)
    if s.startswith("(") and s.endswith(")"):
        sign, s = -sign, s[1:-1]
    try:
        return sign * float(s)
    except ValueError:
        return None


def _is_debit(h):
    return bool(re.search(r"\b(debits?|dr)\b", h))


def _is_credit(h):
    return bool(re.search(r"\b(credits?|cr)\b", h))


def _csv_to_wb(data: bytes) -> Workbook:
    text = data.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for r, row in enumerate(csv.reader(io.StringIO(text), dialect), 1):
        for c, val in enumerate(row, 1):
            if val == "":
                continue
            if PLAIN_NUM_RE.fullmatch(val.strip()):
                ws.cell(r, c).value = parse_amount(val)
            else:
                ws.cell(r, c).value = val
    return wb


def _open_wb(data: bytes, ext: str, values: bool):
    if ext == "csv":
        return _csv_to_wb(data)
    return load_workbook(io.BytesIO(data), data_only=values, keep_vba=(ext == "xlsm"))


@st.cache_resource(show_spinner=False)
def _load_values_wb(sig: str, _data: bytes, ext: str):
    """Read-only use: never mutate the returned workbook."""
    return _open_wb(_data, ext, values=True)


@st.cache_data(show_spinner=False)
def _uncached_formula_count(sig: str, _data: bytes, ext: str, sheet: str,
                            first_row: int, cols: tuple) -> int:
    """Formula cells in the amount columns that have no cached value."""
    if ext not in ("xlsx", "xlsm") or not cols:
        return 0
    lo, hi = min(cols), max(cols)
    wf = load_workbook(io.BytesIO(_data), read_only=True, data_only=False)[sheet]
    wv = load_workbook(io.BytesIO(_data), read_only=True, data_only=True)[sheet]
    bad = 0
    for rf, rv in zip(
        wf.iter_rows(min_row=first_row, min_col=lo, max_col=hi, values_only=True),
        wv.iter_rows(min_row=first_row, min_col=lo, max_col=hi, values_only=True),
    ):
        for f, v in zip(rf, rv):
            if isinstance(f, str) and f.startswith("=") and v is None:
                bad += 1
    return bad


# ───────────────────────────── layout detection ──────────────────────────
def _band_labels(ws, r):
    """Merged band text sitting on the row above the header (e.g. 'Closing Balance')."""
    ctx = {}
    if r > 1:
        for rng in ws.merged_cells.ranges:
            if rng.min_row <= r - 1 <= rng.max_row:
                top = ws.cell(rng.min_row, rng.min_col).value
                if isinstance(top, str) and top.strip():
                    for c in range(rng.min_col, rng.max_col + 1):
                        ctx[c] = top.strip()
    return ctx


def _make_layout(ws, r, name_col, drs, crs, singles):
    ctx = _band_labels(ws, r)

    def hdr(c):
        return str(ws.cell(r, c).value).strip()

    pairs, used = [], set()
    for d in drs:
        nxt = [c for c in crs if c > d and c not in used]
        cr = nxt[0] if nxt else None
        if cr:
            used.add(cr)
        pairs.append((d, cr))
    if not drs and not crs:
        pairs = [(c, None) for c in singles]
    labels = []
    for d, cr in pairs:
        base = hdr(d) + (f" | {hdr(cr)}" if cr else " (signed)")
        band = ctx.get(d, "")
        cols = f"{get_column_letter(d)}" + (f", {get_column_letter(cr)}" if cr else "")
        labels.append(f"{band + ' - ' if band else ''}{base}  [{cols}]")
    return Layout(r, name_col, pairs, labels)


def detect_layout(ws, scan_rows: int = 40):
    """Best-effort guess of header row, name column and amount column pairs."""
    max_c = ws.max_column
    last = min(scan_rows, ws.max_row)

    def row_texts(r):
        out = {}
        for c in range(1, max_c + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip():
                out[c] = v.strip().lower()
        return out

    # pass 1: name header and amount headers on the same row
    for r in range(1, last + 1):
        texts = row_texts(r)
        drs = [c for c, t in texts.items() if _is_debit(t)]
        crs = [c for c, t in texts.items() if _is_credit(t)]
        names = [c for c, t in texts.items()
                 if any(k in t for k in NAME_KEYS) and c not in drs and c not in crs]
        singles = [c for c, t in texts.items()
                   if c not in drs and c not in crs and c not in names
                   and any(k in t for k in SINGLE_AMT_KEYS)]
        if names and (drs or crs or singles):
            return _make_layout(ws, r, names[0], drs, crs, singles)

    # pass 2: Debit/Credit headers on their own row, 'Particulars' a few rows above
    for r in range(1, last + 1):
        texts = row_texts(r)
        drs = [c for c, t in texts.items() if _is_debit(t)]
        crs = [c for c, t in texts.items() if _is_credit(t)]
        if not (drs and crs):
            continue
        name_col = None
        for rr in range(r, max(0, r - 8), -1):
            for c, t in row_texts(rr).items():
                if any(k in t for k in NAME_KEYS) and c not in drs and c not in crs:
                    name_col = c
                    break
            if name_col:
                break
        if name_col is None:
            name_col = 1
        return _make_layout(ws, r, name_col, drs, crs, [])
    return None


# ───────────────────────────── row reading ───────────────────────────────
def read_rows(ws, layout: Layout, pairs: list):
    """Return (rows, total_rows). Trailing 'Total' style rows are split out."""
    rows = []
    for r in range(layout.header_row + 1, ws.max_row + 1):
        raw = ws.cell(r, layout.name_col).value
        raw = "" if raw is None else str(raw)
        dr, cr, net = [], [], []
        for d, c in pairs:
            if c is None:
                v = parse_amount(ws.cell(r, d).value) or 0.0
                dr.append(max(v, 0.0)); cr.append(max(-v, 0.0)); net.append(v)
            else:
                a = parse_amount(ws.cell(r, d).value) or 0.0
                b = parse_amount(ws.cell(r, c).value) or 0.0
                dr.append(a); cr.append(b); net.append(a - b)
        has_amt = any(abs(x) > 0 for x in dr + cr)
        if not raw.strip() and not has_amt:
            continue
        rows.append(Row(
            xl_row=r, name=raw.strip(),
            lvl_space=len(raw) - len(raw.lstrip(" \u00a0\u3000\t")),
            lvl_indent=int(getattr(ws.cell(r, layout.name_col).alignment, "indent", 0) or 0),
            lvl_outline=int(ws.row_dimensions[r].outlineLevel or 0),
            dr=dr, cr=cr, net=net, has_amt=has_amt,
        ))
    totals = []
    while rows and len(totals) < 3 and TOTAL_RE.search(rows[-1].name):
        totals.append(rows.pop())
    return rows, totals[::-1]


# ───────────────────────────── hierarchy logic ───────────────────────────
METHODS = {
    "spaces": "Leading spaces in account name",
    "indent": "Excel cell indent",
    "outline": "Excel row outline level",
    "hybrid": "Amounts, guided by indentation",
    "amount": "Amounts only (no hierarchy in file)",
}


def compute_levels(rows, method):
    key = {"spaces": "lvl_space", "indent": "lvl_indent", "outline": "lvl_outline"}[method]
    levels, prev = [], 0
    for r in rows:
        lvl = getattr(r, key) if r.name else prev   # nameless rows sit beside the row above
        levels.append(lvl)
        prev = lvl
    return levels


def level_source(rows):
    """First indentation source that actually varies in this file (else None)."""
    for m in ("spaces", "indent", "outline"):
        if len(set(compute_levels(rows, m))) >= 2:
            return m
    return None


def method_blocks(rows, method, direction, tol):
    if method in ("spaces", "indent", "outline"):
        return blocks_from_levels(compute_levels(rows, method), direction)
    levels = None
    if method == "hybrid":
        src = level_source(rows)
        levels = compute_levels(rows, src) if src else None
    return blocks_from_amounts([r.net for r in rows], direction, tol, levels)


def auto_select(rows, direction, tol):
    """
    Try every usable method and keep the one whose result balances with the fewest
    unverified subtotals (ties go to the earlier, more structural method).
    Some exports carry indentation that is not a clean hierarchy; this catches them.
    """
    cands = [m for m in ("spaces", "indent", "outline")
             if len(set(compute_levels(rows, m))) >= 2]
    if level_source(rows):
        cands.append("hybrid")
    cands.append("amount")
    P = len(rows[0].net)
    best, best_key = cands[-1], None
    for order, m in enumerate(cands):
        blocks = method_blocks(rows, m, direction, tol)
        info = evaluate_groups(rows, blocks, tol)
        removed = {i for i, v in info.items() if v["status"] == "Matches"}
        gap = max(abs(totals_for(rows, removed, p)[0] - totals_for(rows, removed, p)[1])
                  for p in range(P))
        mism = sum(v["status"] == "Mismatch" for v in info.values())
        key = (gap > max(tol, 0.01), mism, order)
        if best_key is None or key < best_key:
            best, best_key = m, key
    return best


def blocks_from_levels(levels, direction):
    """{group_index: (lo, hi)} - half-open range of rows that sit under the group."""
    n, blocks = len(levels), {}
    if direction == "before":                       # subtotal printed above its ledgers
        for i in range(n - 1):
            if levels[i + 1] > levels[i]:
                j = i + 1
                while j < n and levels[j] > levels[i]:
                    j += 1
                blocks[i] = (i + 1, j)
    else:                                           # subtotal printed below its ledgers
        for i in range(1, n):
            if levels[i - 1] > levels[i]:
                k = i - 1
                while k >= 0 and levels[k] > levels[i]:
                    k -= 1
                blocks[i] = (k + 1, i)
    return blocks


def blocks_from_amounts(nets, direction, tol, levels=None, max_span=5000):
    """
    Infer subtotals from the numbers: a row is a subtotal if the rows that follow it
    (or precede it) add up to its own amount in every selected column pair. Bottom-up,
    so nested subtotals contract before their parents are tested.
    `nets` holds one list of per-pair net amounts per row. When `levels` is given, a
    subtotal's first row beneath it must be indented deeper than the subtotal itself;
    this stops two ledgers with identical balances being read as parent and child.
    """
    n = len(nets)
    seq = nets if direction == "before" else nets[::-1]
    lv = None if levels is None else (levels if direction == "before" else levels[::-1])
    P = len(nets[0]) if n else 0
    nxt = list(range(1, n + 1))
    found = {}
    for p in range(n - 1, -1, -1):
        v = seq[p]
        if all(abs(x) <= tol for x in v):
            continue
        run, u, steps = [0.0] * P, p + 1, 0
        while u < n and steps < max_span:
            for k in range(P):
                run[k] += seq[u][k]
            steps += 1
            end = nxt[u]
            if all(abs(run[k] - v[k]) <= tol for k in range(P)):
                if lv is None or lv[p + 1] > lv[p]:
                    nxt[p] = end
                    found[p] = end
                break
            u = end
    if direction == "before":
        return {p: (p + 1, e) for p, e in found.items()}
    return {n - 1 - p: (n - e, n - 1 - p) for p, e in found.items()}


def evaluate_groups(rows, blocks, tol):
    """Tie each group's amount to the net of the ledgers under it."""
    n = len(rows)
    P = len(rows[0].net) if rows else 0
    pref = [[0.0] * (n + 1) for _ in range(P)]
    for p in range(P):
        s = 0.0
        for i in range(n):
            if i not in blocks:
                s += rows[i].net[p]
            pref[p][i + 1] = s
    out = {}
    for i, (lo, hi) in blocks.items():
        g = rows[i].net
        c = [pref[p][hi] - pref[p][lo] for p in range(P)]
        if not rows[i].has_amt:
            status = "No amount"
        elif all(abs(a - b) <= tol for a, b in zip(g, c)):
            status = "Matches"
        else:
            status = "Mismatch"
        out[i] = {"status": status, "gnet": g, "cnet": c}
    return out


def depth_map(n, blocks):
    diff = [0] * (n + 1)
    for lo, hi in blocks.values():
        diff[lo] += 1
        diff[hi] -= 1
    out, run = [], 0
    for i in range(n):
        run += diff[i]
        out.append(run)
    return out


def totals_for(rows, removed, p):
    dr = sum(r.dr[p] for k, r in enumerate(rows) if k not in removed)
    cr = sum(r.cr[p] for k, r in enumerate(rows) if k not in removed)
    return dr, cr


def parse_row_list(text, valid):
    """'12, 45-50' -> {12,45,..,50} limited to valid Excel rows. Returns (set, ignored)."""
    out, ignored = set(), []
    for part in re.split(r"[,\s;]+", text.strip()):
        if not part:
            continue
        m = re.fullmatch(r"(\d+)(?:-(\d+))?", part)
        if not m:
            ignored.append(part)
            continue
        a, b = int(m.group(1)), int(m.group(2) or m.group(1))
        for x in range(min(a, b), max(a, b) + 1):
            (out.add(x) if x in valid else ignored.append(str(x)))
    return out, ignored


# ───────────────────────────── output workbook ───────────────────────────
def build_output(data, ext, sheet, layout, pairs, pair_labels, rows, total_rows,
                 removed_idx, groups_info, extra_rows, method_label, direction,
                 tol, total_mode, src_name):
    wb = _open_wb(data, ext, values=False)
    ws = wb[sheet]
    cols = [x for pr in pairs for x in pr if x]
    removed_sorted = sorted(removed_idx)

    for i in removed_sorted:
        for c in cols:
            ws.cell(rows[i].xl_row, c).value = None

    if total_rows and total_mode == "rebuild":
        first, last = rows[0].xl_row, rows[-1].xl_row
        t = next((x for x in total_rows if re.match(r"^\s*(grand\s+)?total\b", x.name, re.I)),
                 total_rows[0])
        for c in cols:
            if ws.cell(t.xl_row, c).value is not None:
                L = get_column_letter(c)
                ws.cell(t.xl_row, c).value = f"=SUM({L}{first}:{L}{last})"

    # ---- report sheet
    if REPORT_SHEET in wb.sheetnames:
        del wb[REPORT_SHEET]
    rp = wb.create_sheet(REPORT_SHEET)
    f_norm, f_bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    f_head = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    fill = PatternFill("solid", start_color=ACCENT)
    num = '#,##0.00;(#,##0.00);"-"'

    rp["A1"] = "Trial Balance - Group Subtotal Cleaner: log"
    rp["A1"].font = Font(name="Arial", size=12, bold=True)
    info = [("Source file", src_name), ("Sheet", sheet),
            ("Hierarchy method", method_label),
            ("Subtotal position", "Above its ledgers" if direction == "before" else "Below its ledgers"),
            ("Match tolerance", tol), ("Subtotal rows removed", len(removed_sorted))]
    r = 3
    for k, v in info:
        rp.cell(r, 1, k).font = f_bold
        rp.cell(r, 2, v).font = f_norm
        r += 1

    r += 1
    heads = ["Amount columns", "Debit before", "Credit before", "Difference before",
             "Debit after", "Credit after", "Difference after"]
    for j, h in enumerate(heads, 1):
        cell = rp.cell(r, j, h); cell.font, cell.fill = f_head, fill
    r += 1
    for p, lab in enumerate(pair_labels):
        b_dr, b_cr = totals_for(rows, set(), p)
        a_dr, a_cr = totals_for(rows, set(removed_idx), p)
        vals = [lab, b_dr, b_cr, b_dr - b_cr, a_dr, a_cr, a_dr - a_cr]
        for j, v in enumerate(vals, 1):
            cell = rp.cell(r, j, v); cell.font = f_norm
            if j > 1:
                cell.number_format = num
        r += 1

    r += 1
    P = len(pairs)
    heads = ["Excel row", "Account", "Depth", "Basis"]
    for lab in pair_labels:
        short = lab.split("[")[-1].rstrip("]")
        heads += [f"Removed Dr [{short}]", f"Removed Cr [{short}]", f"Ledgers net [{short}]"]
    for j, h in enumerate(heads, 1):
        cell = rp.cell(r, j, h); cell.font, cell.fill = f_head, fill
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    r += 1
    depth = depth_map(len(rows), {i: g["block"] for i, g in groups_info.items()})
    for i in removed_sorted:
        g = groups_info.get(i)
        basis = (g["status"] if g else "Manual")
        vals = [rows[i].xl_row, rows[i].name, depth[i], basis]
        for p in range(P):
            vals += [rows[i].dr[p], rows[i].cr[p], (round(g["cnet"][p], 2) if g else None)]
        for j, v in enumerate(vals, 1):
            cell = rp.cell(r, j, v); cell.font = f_norm
            if j > 4:
                cell.number_format = num
        r += 1

    rp.column_dimensions["A"].width = 30
    rp.column_dimensions["B"].width = 46
    for j in range(3, len(heads) + 1):
        rp.column_dimensions[get_column_letter(j)].width = 22
    rp.freeze_panes = None

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ───────────────────────────── UI ────────────────────────────────────────
def _stretch():
    try:
        major, minor = (int(x) for x in st.__version__.split(".")[:2])
    except Exception:
        return {"use_container_width": True}
    return {"width": "stretch"} if (major, minor) >= (1, 50) else {"use_container_width": True}


def render_tb_subtotal_cleaner():
    st.markdown("## Trial Balance - Subtotal Cleaner")
    st.caption(
        "Strips group-subtotal amounts from a trial balance so only ledger balances remain "
        "and Debit = Credit. Every subtotal is verified against the ledgers under it."
    )

    up = st.file_uploader("Upload trial balance", type=["xlsx", "xlsm", "csv"], key=NS + "file")
    if up is None:
        st.info("Upload an .xlsx / .xlsm / .csv trial balance to begin.")
        return
    data, fname = up.getvalue(), up.name
    ext = fname.rsplit(".", 1)[-1].lower()
    fsig = hashlib.md5(data).hexdigest()[:10]
    K = lambda s: f"{NS}{fsig}_{s}"          # widget keys are reset per uploaded file

    try:
        wb = _load_values_wb(fsig, data, ext)
    except Exception as e:                    # noqa: BLE001
        st.error(f"Could not open the file: {e}")
        return

    sheet = (st.selectbox("Sheet", wb.sheetnames, key=K("sheet"))
             if len(wb.sheetnames) > 1 else wb.sheetnames[0])
    ws = wb[sheet]
    det = detect_layout(ws)

    # ── layout & options
    with st.expander("Layout and options", expanded=det is None):
        if det is None:
            st.warning("Could not auto-detect the header row. Please set the layout below.")
        c1, c2 = st.columns(2)
        header_row = c1.number_input("Header row (Excel row number)", 1, max(ws.max_row, 1),
                                     det.header_row if det else 1, key=K("hdr"))
        col_letters = [get_column_letter(c) for c in range(1, ws.max_column + 1)]
        name_default = (det.name_col - 1) if det else 0
        name_letter = c2.selectbox("Account-name column", col_letters, index=name_default, key=K("name"))
        name_col = col_letters.index(name_letter) + 1

        pairs, labels = [], []
        manual = st.checkbox("Choose amount columns manually", value=det is None or not det.pairs,
                             key=K("manual"))
        if not manual and det and det.pairs:
            chosen = st.multiselect("Amount columns to clean", det.labels, default=det.labels,
                                    key=K("pairs"))
            for lab in chosen:
                pairs.append(det.pairs[det.labels.index(lab)])
                labels.append(lab)
        else:
            mode = st.radio("Amounts are held in", ["Separate Debit and Credit columns",
                                                     "One signed column (Dr +, Cr -)"],
                            horizontal=True, key=K("mode"))
            m1, m2 = st.columns(2)
            d_letter = m1.selectbox("Debit column" if mode.startswith("Separate") else "Amount column",
                                    col_letters, index=min(3, len(col_letters) - 1), key=K("dcol"))
            d = col_letters.index(d_letter) + 1
            if mode.startswith("Separate"):
                c_letter = m2.selectbox("Credit column", col_letters,
                                        index=min(4, len(col_letters) - 1), key=K("ccol"))
                cc = col_letters.index(c_letter) + 1
                pairs.append((d, cc)); labels.append(f"Debit | Credit  [{d_letter}, {c_letter}]")
            else:
                pairs.append((d, None)); labels.append(f"Signed amount  [{d_letter}]")

        o1, o2, o3 = st.columns(3)
        method_choice = o1.selectbox("Hierarchy source", ["Auto-detect"] + list(METHODS.values()),
                                     key=K("method"))
        pos = o2.radio("Subtotal sits", ["Above its ledgers (Tally style)", "Below its ledgers"],
                       key=K("pos"))
        direction = "before" if pos.startswith("Above") else "after"
        tol = o3.number_input("Match tolerance", 0.0, 1000.0, 0.01, 0.01, format="%.2f", key=K("tol"))

    if not pairs:
        st.warning("Select at least one amount column.")
        return

    layout = Layout(int(header_row), name_col, pairs, labels)
    rows, total_rows = read_rows(ws, layout, pairs)
    if not rows:
        st.error("No data rows found under the header row. Check the layout options.")
        return

    # formulas without cached values would read as blanks
    cols = tuple(x for pr in pairs for x in pr if x)
    bad = _uncached_formula_count(fsig, data, ext, sheet, layout.header_row + 1, cols)
    if bad:
        st.error(f"{bad} amount cells hold formulas with no saved value (file never opened in Excel). "
                 "Open it in Excel, save, and upload again - otherwise those amounts read as blank.")
        return

    # ── hierarchy
    if method_choice == "Auto-detect":
        method = auto_select(rows, direction, tol)
    else:
        method = next(k for k, v in METHODS.items() if v == method_choice)
    if method == "hybrid" and not level_source(rows):
        method = "amount"
        st.info("No usable indentation in this file, so amounts alone are used.")
    blocks = method_blocks(rows, method, direction, tol)
    if method == "amount":
        st.warning("Subtotals are being inferred from the amounts alone. Review the table below "
                   "carefully: two consecutive ledgers with the same balance can be mistaken for "
                   "a parent and child. The Debit/Credit check is your safety net.")
    elif method == "hybrid":
        st.info("This file's indentation is not a clean hierarchy, so subtotals are found from the "
                "amounts (rows that add up to their group) with indentation used only as a guide.")
    info = evaluate_groups(rows, blocks, tol)
    for i in info:
        info[i]["block"] = blocks[i]
    depth = depth_map(len(rows), blocks)

    st.caption(f"Hierarchy read from: **{METHODS[method]}** | rows analysed: **{len(rows)}** | "
               f"total rows set aside: **{len(total_rows)}**")

    # ── review table
    order = {"Mismatch": 0, "Matches": 1, "No amount": 2}
    cand = sorted(info, key=lambda i: (order[info[i]["status"]], rows[i].xl_row))
    P = len(pairs)
    recs = []
    for i in cand:
        rec = {"Remove?": info[i]["status"] == "Matches", "Excel row": rows[i].xl_row,
               "Account": rows[i].name, "Depth": depth[i], "Status": info[i]["status"]}
        for p in range(P):
            sfx = f" [{p + 1}]" if P > 1 else ""
            rec[f"Subtotal net{sfx}"] = round(info[i]["gnet"][p], 2)
            rec[f"Ledgers net{sfx}"] = round(info[i]["cnet"][p], 2)
        recs.append(rec)

    n_match = sum(v["status"] == "Matches" for v in info.values())
    n_mis = sum(v["status"] == "Mismatch" for v in info.values())
    m1, m2, m3, m4 = st.columns(4)
    m1.metric("Subtotal rows found", len(info))
    m2.metric("Verified (tie to ledgers)", n_match)
    m3.metric("Need review", n_mis)
    m4.metric("Ledger rows", len(rows) - len(info))
    if n_mis:
        st.warning("Rows marked Mismatch do not tie to the ledgers under them, so they are NOT pre-selected. "
                   "Tick Remove? only if you are sure they are subtotals.")

    removed_idx = set()
    sel_sig = ""
    if recs:
        import pandas as pd
        df = pd.DataFrame(recs)
        num_cfg = {c: st.column_config.NumberColumn(c, format="%.2f")
                   for c in df.columns if c.startswith(("Subtotal net", "Ledgers net"))}
        edited = st.data_editor(
            df, hide_index=True, key=K(f"ed_{method}_{direction}_{tol}_{len(pairs)}_{layout.header_row}_{name_col}"),
            disabled=[c for c in df.columns if c != "Remove?"],
            column_config={"Remove?": st.column_config.CheckboxColumn("Remove?"), **num_cfg},
            **_stretch(),
        )
        by_row = {rows[i].xl_row: i for i in info}
        for _, rec in edited.iterrows():
            if rec["Remove?"]:
                removed_idx.add(by_row[int(rec["Excel row"])])
    else:
        st.info("No group rows were detected. If subtotals exist, adjust the hierarchy source or "
                "use the manual row list below.")

    extra_txt = st.text_input("Extra rows to treat as subtotals (Excel row numbers, e.g. 12, 40-45)",
                              key=K("extra"))
    extra_rows = set()
    if extra_txt.strip():
        valid = {r.xl_row: k for k, r in enumerate(rows)}
        got, ignored = parse_row_list(extra_txt, valid)
        extra_rows = {valid[x] for x in got}
        removed_idx |= extra_rows
        if ignored:
            st.caption(f"Ignored (not data rows / not understood): {', '.join(ignored[:15])}")

    # ── balance check
    st.markdown("#### Debit / Credit check")
    ok_all = True
    for p, lab in enumerate(labels):
        b_dr, b_cr = totals_for(rows, set(), p)
        a_dr, a_cr = totals_for(rows, removed_idx, p)
        diff = a_dr - a_cr
        ok = abs(diff) <= max(tol, 0.01)
        ok_all &= ok
        st.markdown(f"**{lab}**")
        c1, c2, c3 = st.columns(3)
        c1.metric("Debit (before to after)", f"{a_dr:,.2f}", f"{a_dr - b_dr:,.2f}", delta_color="off")
        c2.metric("Credit (before to after)", f"{a_cr:,.2f}", f"{a_cr - b_cr:,.2f}", delta_color="off")
        c3.metric("Difference", f"{diff:,.2f}", f"was {b_dr - b_cr:,.2f}", delta_color="off")
    if total_rows:
        t0 = next((x for x in total_rows if re.match(r"^\s*(grand\s+)?total\b", x.name, re.I)),
                  total_rows[0])
        for p, lab in enumerate(labels):
            a_dr, a_cr = totals_for(rows, removed_idx, p)
            same = abs(t0.dr[p] - a_dr) <= max(tol, 0.01) and abs(t0.cr[p] - a_cr) <= max(tol, 0.01)
            st.caption(f"File's own total row ({t0.name}): Dr {t0.dr[p]:,.2f} | Cr {t0.cr[p]:,.2f} - "
                       + ("agrees with the cleaned ledger totals." if same
                          else "differs from the cleaned ledger totals (informational)."))
    if ok_all:
        st.success("Debit and credit tally after removing the selected subtotals.")
    else:
        st.error("Debit and credit still do not tally. Check the Mismatch rows above, change the hierarchy "
                 "source or subtotal position, or add missed rows in the box above.")

    # ── total row handling + output
    total_mode = "leave"
    if total_rows:
        tr = ", ".join(str(t.xl_row) for t in total_rows)
        tm = st.radio(f"Total row (Excel row {tr})",
                      ["Rebuild with live SUM formulas", "Leave as it is"], horizontal=True, key=K("tot"))
        total_mode = "rebuild" if tm.startswith("Rebuild") else "leave"

    sig = hashlib.md5(repr((fsig, sheet, layout, method, direction, tol, total_mode,
                            sorted(removed_idx))).encode()).hexdigest()
    if st.button("Build cleaned workbook", type="primary", key=K("build")):
        groups_info = {i: info[i] for i in info}
        out = build_output(data, ext, sheet, layout, pairs, labels, rows, total_rows,
                           removed_idx, groups_info, extra_rows, METHODS[method], direction,
                           tol, total_mode, fname)
        base = fname.rsplit(".", 1)[0]
        out_ext = "xlsm" if ext == "xlsm" else "xlsx"
        st.session_state[NS + "out"] = (sig, out, f"{base}_ledgers_only.{out_ext}")
    saved = st.session_state.get(NS + "out")
    if saved and saved[0] == sig:
        mime = ("application/vnd.ms-excel.sheet.macroEnabled.12" if saved[2].endswith("xlsm")
                else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        st.download_button("Download cleaned workbook", saved[1], file_name=saved[2], mime=mime,
                           key=K("dl"))
        st.caption("Removed amounts are logged on the 'TB Clean Report' sheet. Other formulas in the "
                   "sheet are left exactly as they were.")
    elif saved:
        st.info("Selections changed since the last build - click Build again.")


# ───────────────────────────── standalone runner ─────────────────────────
if __name__ == "__main__":
    st.set_page_config(page_title="TB Subtotal Cleaner", page_icon=":ledger:", layout="wide")
    st.markdown(
        f"""<style>
        .stApp {{ background:#FAF6F1; }}
        .stButton>button[kind="primary"], .stDownloadButton>button {{
            background:#{ACCENT}; border-color:#{ACCENT}; color:#fff; }}
        </style>""",
        unsafe_allow_html=True,
    )
    render_tb_subtotal_cleaner()
